from pathlib import Path
import json, re, html
from docx import Document
from docx.shared import Pt, Inches
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT

root=Path(__file__).resolve().parent.parent
plan=json.loads((root/'data/roadmap.json').read_text(encoding='utf-8'))
index=json.loads((root/'data/notes/index.json').read_text(encoding='utf-8'))
styles=getSampleStyleSheet()
styles['BodyText'].fontSize=10
styles['BodyText'].leading=15
styles['Heading1'].textColor=colors.HexColor('#194969')
styles['Heading2'].textColor=colors.HexColor('#194969')
styles.add(ParagraphStyle('CodeSmall',fontName='Courier',fontSize=7,leading=10))
doc=Document()
doc.styles['Normal'].font.size=Pt(10)
story=[]
def clean(s):
    return s.replace('→',' -> ').replace('←',' <- ').replace('—',' - ').replace('–','-').replace('’',"'").replace('“','"').replace('”','"').replace('·',' / ').replace('≈','about ').replace('✓','yes').replace('Θ','Theta').replace('Ω','Omega')
def add(text,kind='BodyText'):
    story.append(Paragraph(html.escape(clean(text)),styles[kind]))
    if kind=='Title': doc.add_heading(text,0)
    elif kind.startswith('Heading'): doc.add_heading(text,int(kind[-1]))
    else: doc.add_paragraph(text)
def footer(canvas,pdf):
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#667085'))
    canvas.drawString(40,25,'SEBI Grade A IT / independent preparation guide / 12 September 2026')
    canvas.drawRightString(555,25,str(pdf.page))
add('SEBI Grade A IT', 'Title')
add('Your 180-day preparation handbook','Heading1')
add('Four focused hours a day. 26 weeks. Learn, practise, recall and review.')
add('This independent study guide uses the official 2025 recruitment notice as its baseline. The next exam date is not assumed. Verify any newer notice at https://www.sebi.gov.in/sebiweb/about/AboutAction.do?doVacancies=yes')
add('How to study','Heading1')
for text in ['Choose your start date in Settings. The tracker schedules Day 1 from that date.','Open Today’s Targets, follow each linked lesson, and mark a task complete after doing the work.','Use the timer or manually log study minutes. Completing a checkbox does not invent study hours.','Every seventh plan day is review and testing, regardless of calendar weekday.','Use the included original question bank to learn. Supplement it with full-length mocks and the official handout.','Export progress regularly. Import replaces another device’s progress; local document files must be transferred separately.','Core lessons come first. Supplementary computer science and interview references are not claimed to be separate weighted exam topics.']:
    add(text)
add('Weekly targets','Heading1')
for week in plan['weeks']:add(f"Week {week['n']}: {week['theme']}. Days {week['days'][0]['n']}–{week['days'][-1]['n']}.")
for week in plan['weeks']:
    story.append(PageBreak());doc.add_page_break()
    add(f"Week {week['n']} / {week['theme']}",'Heading1')
    add(week['phaseName'])
    for day in week['days']:
        add(f"Day {day['n']} / 240 minutes",'Heading2')
        for task in day['tasks']:add(f"[ ] {task['m']} min / {task['t']}")
        add('Actual minutes: ______   Main mistake: __________________   Next review: ______')
doc.save(root/'library/study-handbook.docx')
SimpleDocTemplate(str(root/'library/study-handbook.pdf'),rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)

# The collection uses the exact in-app notes. Unicode font supports formulas and symbols.
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
font=Path('C:/Windows/Fonts/arial.ttf')
if font.exists():
    pdfmetrics.registerFont(TTFont('StudyFont',str(font)))
    for style in ['BodyText','Heading1','Heading2','Heading3','Title']: styles[style].fontName='StudyFont'
story=[Paragraph('Complete study notes',styles['Title']),Paragraph('Core lessons and supplementary references. Independent study material, not official SEBI content. Policy summaries may become dated: verify linked primary sources before relying on current rules or numerical thresholds.',styles['BodyText'])]
for note in index['notes']:
    story.append(PageBreak());story.append(Paragraph(html.escape(note['title']),styles['Heading1']))
    story.append(Paragraph('Supplementary reference / interview' if note.get('supplemental') else 'Core preparation lesson',styles['BodyText']))
    text=(root/'data/notes'/f"{note['id']}.md").read_text(encoding='utf-8')
    fenced=False
    for line in text.splitlines():
        if line.startswith('```'): fenced=not fenced;continue
        if line.startswith('# '):continue
        if not line.strip():story.append(Spacer(1,5));continue
        if fenced:
            # Wrap very long code lines for the printed collection.
            for start in range(0,len(line),92):story.append(Preformatted(line[start:start+92],styles['CodeSmall']))
        else:
            level=len(line)-len(line.lstrip('#'))
            style='Heading2' if level==2 else 'Heading3' if level>=3 else 'BodyText'
            line=re.sub(r'^#+\s*','',line)
            story.append(Paragraph(html.escape(line),styles[style]))
SimpleDocTemplate(str(root/'library/complete-study-notes.pdf'),rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
print('Created PDF handbook, DOCX handbook and complete study-notes PDF.')
